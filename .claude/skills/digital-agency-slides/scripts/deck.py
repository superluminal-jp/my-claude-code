#!/usr/bin/env python3
"""Render, check and convert a 16:9 HTML deck into an editable PowerPoint file.

The pipeline has one invariant: the deck is authored on a 1280x720 CSS pixel
canvas, which is exactly 13.333x7.5in — PowerPoint's standard 16:9 slide — so
1px = 0.75pt = 9525 EMU and no stage ever rescales geometry. Everything the
converter emits is a native PowerPoint object placed at its measured position;
only content explicitly marked `data-pptx="raster"` becomes a picture.

Subcommands:
    lint      report layout, legibility and accessibility findings
    preview   write one PNG per slide
    pptx      convert the deck to .pptx
    ir        dump the extracted intermediate representation as JSON

Requires: playwright (with chromium) and python-pptx — see scripts/setup.sh.
"""

from __future__ import annotations

import argparse
import base64
import copy
import json
import os
import re
import sys
import urllib.request
import zipfile
from xml.sax.saxutils import quoteattr
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

EMU_PER_PX = 9525  # 914400 EMU/in ÷ 96 px/in — exact
PT_PER_PX = 0.75  # 72 pt/in ÷ 96 px/in — exact
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# "No Style, No Grid": lets the fills and borders read off the page win instead
# of a PowerPoint table style repainting them.
TABLE_STYLE_PLAIN = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

DEFAULTS = {
    "safe_area_px": 32,
    "min_font_px": 14,
    "body_font_px": 18,
    "max_list_items": 5,
    "max_slide_items": 10,
    # The Digital Agency's colour foundation removes WCAG's large-text
    # relaxation: all text clears 4.5:1 whatever its size.
    "text_contrast": 4.5,
    "nontext_contrast": 3.0,
    "slide_selector": ".sld-slide, [data-slide]",
    "secondary_selector": (
        'footer, figcaption, caption, small, [data-role="meta"], '
        ".sld-eyebrow, .sld-card__label, .sld-kpi__label, .sld-steps__index"
    ),
}

# Theme slots filled from DADS tokens so PowerPoint's own colour picker offers
# the design system's palette rather than Office defaults.
# Which design token fills which theme slot. The values are read from the deck
# at conversion time and never written down here: a copy of the palette would
# be a vendored snapshot that drifts silently when the design system moves.
THEME_COLOR_TOKENS = [
    ("dk1", "--color-neutral-solid-gray-900"),
    ("lt1", "--color-neutral-white"),
    ("dk2", "--color-neutral-solid-gray-700"),
    ("lt2", "--color-neutral-solid-gray-50"),
    ("accent1", "--color-key-900"),
    ("accent2", "--color-key-600"),
    ("accent3", "--color-semantic-success-2"),
    ("accent4", "--color-semantic-warning-orange-1"),
    ("accent5", "--color-semantic-error-1"),
    ("accent6", "--color-primitive-purple-700"),
    ("hlink", "--color-primitive-blue-1000"),
    ("folHlink", "--color-primitive-magenta-900"),
]

# The theme is derived from the design system's tokens; it is not the design
# system. Naming it after the design system would misattribute it.
THEME_NAME = "Slide deck (design tokens)"

DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
ADEC = "http://schemas.microsoft.com/office/drawing/2017/decorative"

# OOXML wants a full language tag; the deck declares a subtag.
LANG_ALIASES = {"ja": "ja-JP", "en": "en-US", "zh": "zh-CN", "ko": "ko-KR"}


def px_to_emu(value: float) -> int:
    return int(round(value * EMU_PER_PX))


def px_to_pt(value: float) -> float:
    return round(value * PT_PER_PX, 2)


# --------------------------------------------------------------- extraction --


@dataclass
class Extraction:
    deck: dict[str, Any]
    slides: list[dict[str, Any]]
    findings: list[dict[str, Any]]
    tokens: dict[str, str]
    rasters: dict[str, bytes]
    previews: list[bytes]


def _read_token_css(page) -> dict[str, str]:
    names = [token for _, token in THEME_COLOR_TOKENS] + ["--font-family-sans", "--font-family-mono"]
    return page.evaluate(
        """(names) => {
            const cs = getComputedStyle(document.documentElement);
            const out = {};
            for (const name of names) out[name] = cs.getPropertyValue(name).trim();
            return out;
        }""",
        names,
    )


def extract(
    deck_path: Path,
    options: dict[str, Any],
    *,
    want_rasters: bool,
    want_previews: bool,
) -> Extraction:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:  # pragma: no cover - environment guard
        sys.exit("playwright is not installed. Run scripts/setup.sh first.")

    script = (Path(__file__).parent / "extract_slides.js").read_text(encoding="utf-8")

    launch_kwargs: dict[str, Any] = {}
    executable = options.get("chromium_path") or os.environ.get("DECK_CHROMIUM")
    if executable:
        launch_kwargs["executable_path"] = executable

    with sync_playwright() as pw:
        browser = pw.chromium.launch(**launch_kwargs)
        context = browser.new_context(
            viewport={"width": int(options["viewport_px"]), "height": 720},
            device_scale_factor=2,
        )
        page = context.new_page()
        page.goto(deck_path.resolve().as_uri(), wait_until="load")
        page.evaluate("() => document.fonts.ready")
        page.wait_for_timeout(options["settle_ms"])
        page.add_script_tag(content=script)
        result = page.evaluate("(opts) => window.__extractDeck(opts)", options)

        tokens = _read_token_css(page)
        rasters: dict[str, bytes] = {}
        previews: list[bytes] = []

        if want_rasters:
            for slide in result["slides"]:
                for shape in slide["shapes"]:
                    if shape["kind"] == "raster":
                        locator = page.locator(f'[data-pptx-id="{shape["id"]}"]')
                        rasters[shape["id"]] = locator.screenshot(omit_background=True)

        if want_previews:
            handles = page.query_selector_all(options["slide_selector"])
            previews = [handle.screenshot() for handle in handles]

        browser.close()

    return Extraction(
        deck=result["deck"],
        slides=result["slides"],
        findings=result["findings"],
        tokens=tokens,
        rasters=rasters,
        previews=previews,
    )


# ------------------------------------------------------------------ findings --


def report_findings(findings: Iterable[dict[str, Any]], stream=sys.stderr) -> int:
    findings = sorted(findings, key=lambda f: (f["slide"], f["level"] != "error", f["code"]))
    errors = 0
    for finding in findings:
        if finding["level"] == "error":
            errors += 1
        location = finding.get("element") or ""
        excerpt = (finding.get("text") or "").strip()
        suffix = f"  [{location} {excerpt!r}]" if location else ""
        print(
            f"slide {finding['slide']:>2}  {finding['level']:<5} {finding['code']:<14} {finding['message']}{suffix}",
            file=stream,
        )
    warnings = len(findings) - errors
    print(f"\n{errors} error(s), {warnings} warning(s)", file=stream)
    return errors


# ---------------------------------------------------------------- pptx build --


MAX_IMAGE_BYTES = 32 * 1024 * 1024


def _load_image_bytes(src: str, base: Path) -> bytes | None:
    """Return the bytes, or None with a reason. A broken image reference is an
    ordinary authoring mistake — the browser still lays out a placeholder box —
    so it must degrade to one warning, not end the conversion."""
    import binascii
    from urllib.parse import unquote_to_bytes, urlsplit

    try:
        if src.startswith("data:"):
            head, _, payload = src.partition(",")
            # A data: URI is only base64 when it says so; inline SVG is usually
            # percent-encoded.
            return base64.b64decode(payload) if ";base64" in head else unquote_to_bytes(payload)
        if src.startswith("file://"):
            return Path(urllib.request.url2pathname(urlsplit(src).path)).read_bytes()
        if src.startswith(("http://", "https://")):
            with urllib.request.urlopen(src, timeout=30) as response:  # noqa: S310 - author-supplied deck asset
                blob = response.read(MAX_IMAGE_BYTES + 1)
            if len(blob) > MAX_IMAGE_BYTES:
                print(f"image exceeds {MAX_IMAGE_BYTES} bytes, skipped: {src}", file=sys.stderr)
                return None
            return blob
        candidate = base / src
        return candidate.read_bytes() if candidate.is_file() else None
    except (OSError, ValueError, binascii.Error) as exc:
        print(f"could not load image {src}: {exc}", file=sys.stderr)
        return None


def _sub_element(parent, tag: str, **attrs):
    from lxml import etree

    element = etree.SubElement(parent, f"{{{A}}}{tag}")
    for key, value in attrs.items():
        element.set(key, str(value))
    return element


def _strip_theme_style(shape) -> None:
    """Drop the theme shape style python-pptx attaches to a new autoshape.

    It carries `effectRef`, which paints a drop shadow no slide asked for, and
    a `fontRef` that can override explicit run colours. Everything this
    converter needs is set explicitly on the shape itself.
    """
    element = shape._element
    for style in element.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}style"):
        element.remove(style)
    shape.shadow.inherit = False


def _mark_decorative(shape) -> None:
    """Office's "Mark as decorative", which is the only thing that makes a
    screen reader and the Accessibility Checker skip a shape. An empty `descr`
    does neither — and python-pptx seeds a picture's `descr` with its filename,
    so an unlabelled picture is announced as "image.png"."""
    from lxml import etree

    c_nv_pr = shape._element._nvXxPr.cNvPr
    c_nv_pr.set("descr", "")
    ext_lst = c_nv_pr.find(f"{{{A}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(c_nv_pr, f"{{{A}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{A}}}ext")
    ext.set("uri", DECORATIVE_URI)
    etree.SubElement(ext, f"{{{ADEC}}}decorative", nsmap={"adec": ADEC}).set("val", "1")


def _set_alt_text(shape, text: str) -> None:
    if text:
        shape._element._nvXxPr.cNvPr.set("descr", text)
    else:
        _mark_decorative(shape)


def _ooxml_lang(tag: str | None) -> str:
    tag = (tag or "").strip()
    if not tag:
        return "ja-JP"
    return LANG_ALIASES.get(tag.lower(), tag if "-" in tag else tag)


def _apply_lang(run, lang: str) -> None:
    """OOXML carries language per run. python-pptx writes none, so every run
    inherits `lang="en-US"` from the presentation's default text style and a
    screen reader reads Japanese with an English voice."""
    r_pr = run._r.get_or_add_rPr()
    r_pr.set("lang", lang)
    if lang.lower().startswith(("ja", "zh", "ko")):
        r_pr.set("altLang", "en-US")


def _apply_east_asian_font(run, name: str) -> None:
    """python-pptx only writes the latin typeface; Japanese text needs `ea`."""
    r_pr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        existing = r_pr.find(f"{{{A}}}{tag}")
        if existing is not None:
            r_pr.remove(existing)
    _sub_element(r_pr, "ea", typeface=name)
    _sub_element(r_pr, "cs", typeface=name)


def _apply_letter_spacing(run, spacing_px: float) -> None:
    if abs(spacing_px) < 0.01:
        return
    run._r.get_or_add_rPr().set("spc", str(int(round(spacing_px * PT_PER_PX * 100))))


def _apply_bullet(paragraph, marker: dict[str, Any] | None, mar_l_px: float, indent_px: float, font: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(px_to_emu(max(0.0, mar_l_px))))
    p_pr.set("indent", str(px_to_emu(indent_px)))
    for tag in ("buNone", "buChar", "buAutoNum", "buFont"):
        for existing in p_pr.findall(f"{{{A}}}{tag}"):
            p_pr.remove(existing)
    if not marker or marker.get("kind") == "none":
        _sub_element(p_pr, "buNone")
        return
    _sub_element(p_pr, "buFont", typeface=font)
    if marker["kind"] == "number":
        _sub_element(p_pr, "buAutoNum", type="arabicPeriod")
    else:
        _sub_element(p_pr, "buChar", char=marker.get("char", "•"))


def _alignment(value: str):
    from pptx.enum.text import PP_ALIGN

    return {
        "left": PP_ALIGN.LEFT,
        "start": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "end": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get(value, PP_ALIGN.LEFT)


def _anchor(value: str):
    from pptx.enum.text import MSO_ANCHOR

    return {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(value, MSO_ANCHOR.TOP)


def _font_for(name: str, font_map: dict[str, str]) -> str:
    return font_map.get(name, font_map.get("*", name))


def _slack_rect(rect: dict[str, float], align: str, slack: float) -> tuple[float, float]:
    """Widen a wrapping text box away from the edge its text is anchored to.

    The box is measured against the browser's font metrics. PowerPoint resolves
    the same font slightly differently, and a box sized to the exact measured
    width turns that difference into an unwanted line break. The extra width
    goes where it cannot move the visible text.
    """
    if slack <= 0:
        return rect["x"], rect["w"]
    if align in ("center",):
        return rect["x"] - slack, rect["w"] + 2 * slack
    if align in ("right", "end"):
        return rect["x"] - slack, rect["w"] + slack
    return rect["x"], rect["w"] + slack


def _add_text_shape(
    slide,
    shape_ir: dict[str, Any],
    font_map: dict[str, str],
    lang: str,
    slack: float = 0.0,
    wrap_policy: str = "preserve",
) -> None:
    from pptx.util import Emu

    rect = shape_ir["rect"]
    wrap = True if wrap_policy == "reflow" else shape_ir.get("wrap", True)
    align = shape_ir["paragraphs"][0]["align"] if shape_ir["paragraphs"] else "left"
    x, w = _slack_rect(rect, align, slack if wrap else 0.0)
    box = slide.shapes.add_textbox(
        Emu(px_to_emu(x)), Emu(px_to_emu(rect["y"])), Emu(px_to_emu(w)), Emu(px_to_emu(rect["h"]))
    )
    _fill_text_frame(box, shape_ir, font_map, lang, slack, wrap_policy)


def _fill_text_frame(
    box,
    shape_ir: dict[str, Any],
    font_map: dict[str, str],
    lang: str,
    slack: float = 0.0,
    wrap_policy: str = "preserve",
) -> None:
    """Place and fill an existing shape. Used for ordinary text boxes and for
    the title placeholder, so a promoted heading is formatted identically.

    `wrap_policy` is the delivery-versus-editing trade. "preserve" keeps the
    browser's line breaking, so a label the browser fit on one line cannot be
    re-wrapped by PowerPoint's own font metrics — right when the file is only
    presented. "reflow" lets every box wrap, so text edited in PowerPoint stays
    inside its box instead of running off the slide — right when somebody will
    retype the content there."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    rect = shape_ir["rect"]
    wrap = True if wrap_policy == "reflow" else shape_ir.get("wrap", True)
    align = shape_ir["paragraphs"][0]["align"] if shape_ir["paragraphs"] else "left"
    x, w = _slack_rect(rect, align, slack if wrap else 0.0)
    box.left, box.top = Emu(px_to_emu(x)), Emu(px_to_emu(rect["y"]))
    box.width, box.height = Emu(px_to_emu(w)), Emu(px_to_emu(rect["h"]))
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = _anchor(shape_ir.get("anchor", "top"))

    while len(frame.paragraphs) > 1:
        frame._txBody.remove(frame.paragraphs[-1]._p)
    for index, para_ir in enumerate(shape_ir["paragraphs"]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = _alignment(para_ir["align"])
        if para_ir.get("line_px"):
            paragraph.line_spacing = Pt(px_to_pt(para_ir["line_px"]))
        paragraph.space_before = Pt(px_to_pt(para_ir.get("space_before_px", 0)))
        paragraph.space_after = Pt(px_to_pt(para_ir.get("space_after_px", 0)))

        first_font = _font_for(para_ir["runs"][0]["font"], font_map) if para_ir["runs"] else "Arial"
        marker = para_ir.get("marker")
        if marker:
            # A bullet hangs left of the text: the margin is where the glyphs
            # start, the indent is the negative distance back to the marker.
            text_left = para_ir.get("indent_px", 0.0)
            box_left = para_ir.get("box_left_px", text_left)
            _apply_bullet(paragraph, marker, text_left, box_left - text_left, first_font)
        else:
            # Without a bullet the left margin is the paragraph box's own
            # offset. Using the glyph offset would add the centring or
            # right-alignment gap on top of the alignment itself.
            _apply_bullet(paragraph, None, para_ir.get("box_left_px", 0.0), 0.0, first_font)

        for run_ir in para_ir["runs"]:
            run = paragraph.add_run()
            run.text = run_ir["text"]
            font = run.font
            font.size = Pt(px_to_pt(run_ir["size_px"]))
            font.bold = run_ir["bold"]
            font.italic = run_ir["italic"]
            font.underline = run_ir["underline"]
            if run_ir.get("strike"):
                run._r.get_or_add_rPr().set("strike", "sngStrike")
            if run_ir["color"]:
                font.color.rgb = RGBColor.from_string(run_ir["color"])
            name = _font_for(run_ir["font"], font_map)
            font.name = name
            _apply_east_asian_font(run, name)
            _apply_lang(run, lang)
            _apply_letter_spacing(run, run_ir.get("letter_spacing_px", 0.0))


def _add_decoration(slide, shape_ir: dict[str, Any]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    rect = shape_ir["rect"]
    borders = shape_ir.get("borders", [])
    radius = shape_ir.get("radius", 0) or 0
    uniform = (
        len(borders) == 4
        and len({(b["width"], b["color"]["hex"]) for b in borders}) == 1
    )

    if shape_ir.get("fill") or uniform or radius:
        auto_shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        box = slide.shapes.add_shape(
            auto_shape,
            Emu(px_to_emu(rect["x"])),
            Emu(px_to_emu(rect["y"])),
            Emu(px_to_emu(rect["w"])),
            Emu(px_to_emu(rect["h"])),
        )
        _strip_theme_style(box)
        _mark_decorative(box)
        if radius > 0 and min(rect["w"], rect["h"]) > 0:
            box.adjustments[0] = min(0.5, radius / min(rect["w"], rect["h"]))
        if shape_ir.get("fill"):
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor.from_string(shape_ir["fill"])
            _apply_fill_alpha(box.fill, shape_ir.get("fill_alpha", 1.0))
        else:
            box.fill.background()
        if uniform:
            border = borders[0]
            box.line.color.rgb = RGBColor.from_string(border["color"]["hex"])
            box.line.width = Emu(px_to_emu(border["width"]))
        else:
            box.line.fill.background()
        box.text_frame.text = ""

    if uniform:
        return

    # Per-side borders (a left accent rule, a top keyline) have no PowerPoint
    # equivalent on one shape, so each visible side becomes its own filled bar.
    for border in borders:
        side, width = border["side"], border["width"]
        if side == "top":
            geom = (rect["x"], rect["y"], rect["w"], width)
        elif side == "bottom":
            geom = (rect["x"], rect["y"] + rect["h"] - width, rect["w"], width)
        elif side == "left":
            geom = (rect["x"], rect["y"], width, rect["h"])
        else:
            geom = (rect["x"] + rect["w"] - width, rect["y"], width, rect["h"])
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, *(Emu(px_to_emu(v)) for v in geom)
        )
        _strip_theme_style(bar)
        _mark_decorative(bar)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(border["color"]["hex"])
        bar.line.fill.background()
        bar.text_frame.text = ""


def _apply_fill_alpha(fill, alpha: float) -> None:
    """A CSS colour can be translucent; DrawingML expresses that as an <a:alpha>
    child of the colour. Without it a 5% wash paints as a solid block, which is
    what a design system's opacity-gray tokens would have produced."""
    if alpha >= 0.999:
        return
    srgb = fill.fore_color._xFill.find(f"{{{A}}}srgbClr")
    if srgb is None:
        return
    for existing in srgb.findall(f"{{{A}}}alpha"):
        srgb.remove(existing)
    _sub_element(srgb, "alpha", val=int(round(max(0.0, alpha) * 100000)))


def _set_cell_borders(cell, borders: list[dict[str, Any]]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tags = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}
    for border in borders:
        tag = tags[border["side"]]
        for existing in tc_pr.findall(f"{{{A}}}{tag}"):
            tc_pr.remove(existing)
        line = _sub_element(tc_pr, tag, w=px_to_emu(border["width"]), cap="flat", cmpd="sng", algn="ctr")
        fill = _sub_element(line, "solidFill")
        _sub_element(fill, "srgbClr", val=border["color"]["hex"])


def _add_table(slide, shape_ir: dict[str, Any], font_map: dict[str, str], lang: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    rect = shape_ir["rect"]
    rows_ir = shape_ir["rows"]
    n_rows = len(rows_ir)
    n_cols = len(shape_ir["col_widths_px"])
    frame = slide.shapes.add_table(
        n_rows, n_cols, Emu(px_to_emu(rect["x"])), Emu(px_to_emu(rect["y"])), Emu(px_to_emu(rect["w"])), Emu(px_to_emu(rect["h"]))
    )
    table = frame.table
    # `firstRow` is the only header-row semantic PowerPoint exposes, and without
    # it a screen reader reading a cell never announces its column. It cannot
    # repaint anything here: every cell carries its own fill and borders, and
    # direct cell formatting overrides table-style formatting. Banding is
    # separate, and stays off.
    table.first_row = bool(rows_ir) and all(cell["header"] for cell in rows_ir[0])
    table.horz_banding = False
    tbl_pr = table._tbl.find(f"{{{A}}}tblPr")
    if tbl_pr is not None:
        style_id = tbl_pr.find(f"{{{A}}}tableStyleId")
        if style_id is None:
            style_id = _sub_element(tbl_pr, "tableStyleId")
        style_id.text = TABLE_STYLE_PLAIN

    for index, width in enumerate(shape_ir["col_widths_px"]):
        table.columns[index].width = Emu(px_to_emu(width))
    for index, height in enumerate(shape_ir["row_heights_px"]):
        table.rows[index].height = Emu(px_to_emu(height))

    occupied: set[tuple[int, int]] = set()
    for r, row_ir in enumerate(rows_ir):
        col = 0
        for cell_ir in row_ir:
            while (r, col) in occupied:
                col += 1
            if col >= n_cols:
                break
            cell = table.cell(r, col)
            span_r = min(cell_ir["rowspan"], n_rows - r)
            span_c = min(cell_ir["colspan"], n_cols - col)
            if span_r > 1 or span_c > 1:
                cell.merge(table.cell(r + span_r - 1, col + span_c - 1))
            for dr in range(span_r):
                for dc in range(span_c):
                    occupied.add((r + dr, col + dc))

            cell.margin_left = Emu(px_to_emu(cell_ir["pad"]["l"]))
            cell.margin_right = Emu(px_to_emu(cell_ir["pad"]["r"]))
            cell.margin_top = Emu(px_to_emu(cell_ir["pad"]["t"]))
            cell.margin_bottom = Emu(px_to_emu(cell_ir["pad"]["b"]))
            if cell_ir["fill"]:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(cell_ir["fill"])
            else:
                cell.fill.background()
            _set_cell_borders(cell, cell_ir["borders"])

            frame_ = cell.text_frame
            frame_.word_wrap = cell_ir.get("lines", 1) > 1
            paragraph = frame_.paragraphs[0]
            paragraph.alignment = _alignment(cell_ir["align"])
            fallback = {
                "text": cell_ir["text"],
                "size_px": 16,
                "bold": cell_ir["header"],
                "italic": False,
                "underline": False,
                "strike": False,
                "color": None,
                "font": "sans-serif",
                "letter_spacing_px": 0,
            }
            for run_ir in cell_ir["runs"] or [fallback]:
                run = paragraph.add_run()
                run.text = run_ir["text"]
                run.font.size = Pt(px_to_pt(run_ir["size_px"]))
                run.font.bold = run_ir["bold"]
                run.font.italic = run_ir["italic"]
                if run_ir.get("strike"):
                    run._r.get_or_add_rPr().set("strike", "sngStrike")
                if run_ir["color"]:
                    run.font.color.rgb = RGBColor.from_string(run_ir["color"])
                name = _font_for(run_ir["font"], font_map)
                run.font.name = name
                _apply_east_asian_font(run, name)
                _apply_lang(run, lang)
            col += span_c


def _add_picture(slide, shape_ir: dict[str, Any], blob: bytes) -> None:
    from pptx.util import Emu

    rect = shape_ir["rect"]
    picture = slide.shapes.add_picture(
        BytesIO(blob),
        Emu(px_to_emu(rect["x"])),
        Emu(px_to_emu(rect["y"])),
        Emu(px_to_emu(rect["w"])),
        Emu(px_to_emu(rect["h"])),
    )
    _set_alt_text(picture, shape_ir.get("alt", ""))


def _find_heading_shape(shapes: list[dict[str, Any]], heading: str) -> dict[str, Any] | None:
    """The slide's governing message is already a text shape. Promote that shape
    into the title placeholder rather than adding a second copy, which a screen
    reader would announce twice."""
    if not heading:
        return None
    for shape in shapes:
        if shape["kind"] != "text":
            continue
        text = "".join(run["text"] for para in shape["paragraphs"] for run in para["runs"]).strip()
        if text == heading:
            return shape
    return None


def _apply_slide_title(
    slide,
    slide_ir: dict[str, Any],
    heading: str,
    promoted: dict[str, Any] | None,
    font_map: dict[str, str],
    lang: str,
    slack_px: float,
    wrap_policy: str = "preserve",
) -> None:
    from pptx.util import Emu

    placeholder = slide.shapes.title
    if placeholder is None:
        return
    if promoted is not None:
        _fill_text_frame(placeholder, promoted, font_map, lang, slack_px, wrap_policy)
        return
    # The heading is split across shapes, or absent. Keep the slide navigable
    # with a title parked outside the canvas rather than leaving it untitled.
    placeholder.left, placeholder.top = Emu(px_to_emu(-4000)), Emu(0)
    placeholder.width, placeholder.height = Emu(px_to_emu(1000)), Emu(px_to_emu(60))
    frame = placeholder.text_frame
    frame.text = heading or f"スライド {slide_ir['index']}"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            _apply_lang(run, lang)


def build_pptx(
    extraction: Extraction,
    deck_path: Path,
    out_path: Path,
    font_map: dict[str, str],
    slack_px: float = 0.0,
    wrap_policy: str = "preserve",
) -> list[str]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu

    warnings: list[str] = []
    presentation = Presentation()
    presentation.slide_width = Emu(px_to_emu(extraction.deck["width_px"]))
    presentation.slide_height = Emu(px_to_emu(extraction.deck["height_px"]))
    lang = _ooxml_lang(extraction.deck.get("lang"))
    core = presentation.core_properties
    core.title = (extraction.deck.get("title") or "").strip()
    core.language = lang
    core.last_modified_by = ""
    core.comments = ""

    # "Title Only" is the leanest layout that still carries a title placeholder,
    # and python-pptx clones only title/body/object placeholders, so nothing
    # else arrives on the slide. A slide without one is untitled in the outline,
    # in the reading-order pane and to every screen reader.
    title_layout = presentation.slide_layouts[5]

    for slide_ir in extraction.slides:
        slide = presentation.slides.add_slide(title_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(slide_ir["background"])

        heading = (slide_ir.get("title") or "").strip()
        promoted = _find_heading_shape(slide_ir["shapes"], heading)
        _apply_slide_title(slide, slide_ir, heading, promoted, font_map, lang, slack_px, wrap_policy)

        for shape_ir in slide_ir["shapes"]:
            if shape_ir is promoted:
                continue
            kind = shape_ir["kind"]
            if kind == "text":
                _add_text_shape(slide, shape_ir, font_map, lang, slack_px, wrap_policy)
            elif kind == "shape":
                _add_decoration(slide, shape_ir)
            elif kind == "table":
                _add_table(slide, shape_ir, font_map, lang)
            elif kind == "raster":
                blob = extraction.rasters.get(shape_ir["id"])
                if blob:
                    _add_picture(slide, shape_ir, blob)
                else:
                    warnings.append(f"slide {slide_ir['index']}: raster {shape_ir['id']} was not captured")
            elif kind == "image":
                blob = _load_image_bytes(shape_ir["src"], deck_path.parent)
                if blob:
                    _add_picture(slide, shape_ir, blob)
                else:
                    warnings.append(f"slide {slide_ir['index']}: could not load image {shape_ir['src']}")

        if slide_ir["notes"]:
            notes_frame = slide.notes_slide.notes_text_frame
            notes_frame.text = slide_ir["notes"]
            for paragraph in notes_frame.paragraphs:
                for run in paragraph.runs:
                    _apply_lang(run, lang)

    presentation.save(out_path)
    warnings.extend(_patch_package(out_path, extraction.tokens, font_map, lang))
    return warnings


# ------------------------------------------------------------------- theming --


def _theme_xml(
    tokens: dict[str, str], font_map: dict[str, str]
) -> tuple[str | None, str, list[str]]:
    """Build the theme from the tokens the deck actually resolved.

    A slot whose token cannot be read is not guessed at: the whole colour scheme
    is left alone and the caller is told which token was missing. Substituting a
    remembered value would put a stale copy of someone else's palette in the
    file under this skill's name."""
    warnings: list[str] = []
    values: list[tuple[str, str]] = []
    for slot, token in THEME_COLOR_TOKENS:
        value = (tokens.get(token) or "").strip().lstrip("#")
        if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
            warnings.append(f"theme: {token} did not resolve in the deck; keeping the default colour scheme")
            values = []
            break
        values.append((slot, value.upper()))

    clr_scheme = None
    if values:
        body = "".join(f"<a:{slot}><a:srgbClr val='{value}'/></a:{slot}>" for slot, value in values)
        clr_scheme = f"<a:clrScheme xmlns:a='{A}' name='{THEME_NAME}'>{body}</a:clrScheme>"

    sans = (tokens.get("--font-family-sans") or "").split(",")[0].strip().strip("'\"")
    sans = _font_for(sans, font_map) if sans else ""
    if not sans:
        warnings.append("theme: no sans font resolved in the deck; keeping the default font scheme")
        return clr_scheme, "", warnings
    quoted = quoteattr(sans)
    fonts = "".join(
        f"<a:{scheme}><a:latin typeface={quoted}/><a:ea typeface={quoted}/><a:cs typeface=''/></a:{scheme}>"
        for scheme in ("majorFont", "minorFont")
    )
    font_scheme = f"<a:fontScheme xmlns:a='{A}' name='{THEME_NAME}'>{fonts}</a:fontScheme>"
    return clr_scheme, font_scheme, warnings


def _patch_package(
    pptx_path: Path, tokens: dict[str, str], font_map: dict[str, str], lang: str
) -> list[str]:
    """Rewrite the parts python-pptx cannot reach through its object model:

    the theme, so PowerPoint's colour and font pickers offer the deck's own
    palette; the inherited `lang="en-US"` in the master, the layouts and the
    presentation defaults, so text the recipient types later is also tagged
    correctly; and the slide-size token, which is stale in python-pptx's
    template."""
    clr_scheme, font_scheme, warnings = _theme_xml(tokens, font_map)
    with zipfile.ZipFile(pptx_path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]

    def rewrite_theme(text: str) -> str:
        if clr_scheme:
            text = re.sub(r"<a:clrScheme\b.*?</a:clrScheme>", lambda _: clr_scheme, text, count=1, flags=re.S)
        if font_scheme:
            text = re.sub(r"<a:fontScheme\b.*?</a:fontScheme>", lambda _: font_scheme, text, count=1, flags=re.S)
        return text

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item, data in entries:
            name = item.filename
            if name.endswith(".xml"):
                text = data.decode("utf-8")
                if name.startswith("ppt/theme/"):
                    text = rewrite_theme(text)
                if name in ("ppt/presentation.xml", "ppt/notesMasters/notesMaster1.xml") or name.startswith(
                    ("ppt/slideMasters/", "ppt/slideLayouts/")
                ):
                    text = text.replace('lang="en-US"', f'lang="{lang}"')
                if name == "ppt/presentation.xml":
                    text = re.sub(r'(<p:sldSz[^>]*?)\s*type="[^"]*"', r"\1", text)
                    text = text.replace("<p:sldSz ", '<p:sldSz type="screen16x9" ', 1)
                data = text.encode("utf-8")
            target.writestr(copy.copy(item), data)
    pptx_path.write_bytes(buffer.getvalue())
    return warnings


# ----------------------------------------------------------------------- cli --


def _options(args: argparse.Namespace) -> dict[str, Any]:
    options = dict(DEFAULTS)
    options.update(
        {
            "safe_area_px": args.safe_area,
            "min_font_px": args.min_font,
            "body_font_px": args.body_font,
            "max_list_items": args.max_list_items,
            "max_slide_items": args.max_slide_items,
            "text_contrast": args.text_contrast,
            "nontext_contrast": args.nontext_contrast,
            "slide_selector": args.selector,
            "secondary_selector": args.secondary_selector,
            "viewport_px": args.viewport,
            "settle_ms": args.settle_ms,
            "chromium_path": args.chromium,
        }
    )
    return options


def _font_map(values: list[str] | None) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for value in values or []:
        source, _, target = value.partition("=")
        if not target:
            mapping["*"] = source.strip()
        else:
            mapping[source.strip()] = target.strip()
    return mapping


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["lint", "preview", "pptx", "ir"])
    parser.add_argument("deck", type=Path, help="path to the deck HTML file")
    parser.add_argument("-o", "--out", type=Path, help="output file (pptx/ir) or directory (preview)")
    parser.add_argument("--selector", default=DEFAULTS["slide_selector"], help="CSS selector matching one slide")
    parser.add_argument(
        "--secondary-selector",
        default=DEFAULTS["secondary_selector"],
        help="CSS selector for content exempt from the body-size warning",
    )
    parser.add_argument("--safe-area", type=float, default=DEFAULTS["safe_area_px"], help="safe-area inset in px")
    parser.add_argument("--min-font", type=float, default=DEFAULTS["min_font_px"], help="minimum font size in px")
    parser.add_argument("--body-font", type=float, default=DEFAULTS["body_font_px"], help="comfortable body size in px")
    parser.add_argument("--max-list-items", type=int, default=DEFAULTS["max_list_items"], help="items in any one list")
    parser.add_argument(
        "--max-slide-items", type=int, default=DEFAULTS["max_slide_items"], help="list items on a slide, pooled"
    )
    parser.add_argument(
        "--text-contrast", type=float, default=DEFAULTS["text_contrast"], help="minimum contrast ratio for text"
    )
    parser.add_argument(
        "--nontext-contrast",
        type=float,
        default=DEFAULTS["nontext_contrast"],
        help="minimum contrast ratio for borders and rules",
    )
    parser.add_argument("--viewport", type=int, default=1440, help="browser viewport width in px")
    parser.add_argument("--settle-ms", type=int, default=250, help="wait after load before measuring")
    parser.add_argument(
        "--chromium",
        default=None,
        metavar="PATH",
        help="Chromium executable to measure with (default: Playwright's own; env DECK_CHROMIUM)",
    )
    parser.add_argument(
        "--font",
        action="append",
        metavar="FROM=TO",
        help="remap a font for the .pptx (repeatable); a bare name remaps every font",
    )
    parser.add_argument(
        "--slack",
        type=float,
        default=6.0,
        help="extra px given to wrapping text boxes to absorb font-metric differences",
    )
    parser.add_argument(
        "--wrap",
        choices=["preserve", "reflow"],
        default="preserve",
        help=(
            "preserve: keep the browser's line breaking, so the delivered file looks exactly like the "
            "preview (default). reflow: let every text box wrap, so text edited in PowerPoint stays "
            "inside its box"
        ),
    )
    parser.add_argument("--allow-findings", action="store_true", help="convert even when the deck has errors")
    args = parser.parse_args(argv)

    if not args.deck.exists():
        sys.exit(f"no such deck: {args.deck}")

    options = _options(args)
    wants_pptx = args.command == "pptx"
    extraction = extract(
        args.deck,
        options,
        want_rasters=wants_pptx,
        want_previews=args.command == "preview",
    )

    if args.command == "ir":
        payload = {"deck": extraction.deck, "slides": extraction.slides, "findings": extraction.findings}
        text = json.dumps(payload, ensure_ascii=False, indent=2)
        if args.out:
            args.out.write_text(text, encoding="utf-8")
            print(f"wrote {args.out}")
        else:
            print(text)
        return 0

    if extraction.deck is None:
        report_findings(extraction.findings)
        sys.exit(f"no slide matched {options['slide_selector']}")

    errors = report_findings(extraction.findings)

    if args.command == "lint":
        return 1 if errors else 0

    if args.command == "preview":
        out_dir = args.out or args.deck.parent / "preview"
        out_dir.mkdir(parents=True, exist_ok=True)
        for index, blob in enumerate(extraction.previews, start=1):
            (out_dir / f"slide-{index:02d}.png").write_bytes(blob)
        print(f"wrote {len(extraction.previews)} PNG(s) to {out_dir}")
        return 1 if errors else 0

    if errors and not args.allow_findings:
        print("\nRefusing to convert a deck with errors. Fix them, or pass --allow-findings.", file=sys.stderr)
        return 1

    out_path = args.out or args.deck.with_suffix(".pptx")
    warnings = build_pptx(extraction, args.deck, out_path, _font_map(args.font), args.slack, args.wrap)
    for warning in warnings:
        print(warning, file=sys.stderr)
    shapes = sum(len(slide["shapes"]) for slide in extraction.slides)
    rasters = sum(1 for slide in extraction.slides for shape in slide["shapes"] if shape["kind"] == "raster")
    print(
        f"wrote {out_path}: {len(extraction.slides)} slide(s), {shapes} shape(s), "
        f"{shapes - rasters} native, {rasters} rasterized"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
